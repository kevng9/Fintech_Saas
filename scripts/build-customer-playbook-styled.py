"""Restyle the customer playbook onto the SingleStore presentation template.

Keeps the 19-slide story. Applies the template style guide: 13.33 x 7.5,
Roboto, the brand palette, left alignment, and the template logo and footer.
"""

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path("/Users/kevinng/Downloads/SingleStore Company Presentation Template.pptx")
SOURCE = ROOT / "Fintech_Aura_Analyst_Playbook_Customer.pptx"
DEST = ROOT / "Fintech_Aura_Analyst_Playbook_Customer.pptx"

INK = RGBColor(0x09, 0x09, 0x09)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PURPLE = RGBColor(0x82, 0x0D, 0xDF)
LILAC = RGBColor(0xD1, 0x99, 0xFF)
DEEP = RGBColor(0x36, 0x00, 0x61)
GRAY = RGBColor(0x52, 0x52, 0x52)
BLUE = RGBColor(0x4B, 0x47, 0xFF)
GOLD = RGBColor(0xFF, 0xCB, 0x2E)
ORANGE = RGBColor(0xFF, 0x6C, 0x02)
MIST = RGBColor(0xF2, 0xF2, 0xF2)
LINE = RGBColor(0xD9, 0xD9, 0xD9)

COVER = 37
LIGHT = 44
DARK = 50


def _set_run(run, text, font, size, color):
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = False
    run.font.italic = False
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        node = rpr.find(qn(f"a:{tag}"))
        if node is None:
            node = rpr.makeelement(qn(f"a:{tag}"), {})
            rpr.append(node)
        node.set("typeface", font)


def add_text(slide, x, y, w, h, text, font="Roboto", size=14, color=INK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.paragraphs[0].alignment = align
    anchor_map = {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}
    tf._txBody.bodyPr.set("anchor", anchor_map[anchor])
    lines = text.split("\n") if isinstance(text, str) else list(text)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_before = Pt(0)
        p.space_after = Pt(4 if len(lines) > 1 else 0)
        _set_run(p.add_run(), line, font, size, color)
    return box


def add_rect(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    return shape


def hline(slide, x, y, w, color=LINE):
    return add_rect(slide, x, y, w, 0.01, color)


def style_title(shape, text, size, color):
    shape.left = Inches(0.64)
    shape.top = Inches(0.40)
    shape.width = Inches(9.0)
    shape.height = Inches(0.95)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.clear()
    p.alignment = PP_ALIGN.LEFT
    p.line_spacing = 1.0
    _set_run(p.add_run(), text, "Roboto Medium", size, color)


def light(prs, title, size=24):
    slide = prs.slides.add_slide(prs.slide_layouts[LIGHT])
    style_title(slide.shapes.title, title, size, INK)
    return slide


def dark(prs, title, size=24):
    slide = prs.slides.add_slide(prs.slide_layouts[DARK])
    add_text(slide, 0.64, 0.40, 12.05, 0.95, title, "Roboto Medium", size, WHITE)
    return slide


def label(slide, x, y, w, text, color=GRAY):
    add_text(slide, x, y, w, 0.22, text, "Roboto Medium", 11, color)


def bullets(slide, x, y, w, h, items, size=12, color=INK):
    add_text(slide, x, y, w, h, [f"•  {item}" for item in items], "Roboto", size, color)


def copy_note(slide, source_slide):
    if not source_slide.has_notes_slide:
        return
    text = source_slide.notes_slide.notes_text_frame.text.strip()
    if text:
        slide.notes_slide.notes_text_frame.text = text


def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[COVER])
    title = slide.shapes.title
    title.left = Inches(0.64)
    title.top = Inches(1.70)
    title.width = Inches(11.4)
    title.height = Inches(1.35)
    tf = title.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.clear()
    p.line_spacing = 1.0
    _set_run(p.add_run(), "Our shared goal: make one costly decision much faster",
             "Roboto Medium", 32, LILAC)
    body = slide.placeholders[1]
    body.text_frame.paragraphs[0].clear()
    _set_run(body.text_frame.paragraphs[0].add_run(),
             "Reducing time-to-know is the central promise.", "Roboto", 14, WHITE)
    bullets(slide, 0.64, 3.30, 6.15, 2.9, [
        "Start with a consequential decision, how long it takes to know today, and the cost of delay.",
        "Agree on the observable delta you want before choosing capabilities or screens.",
        "Balance business outcome, regulated risk, and technical feasibility from the start.",
        "Treat every business delta as a hypothesis until your owners validate it.",
    ], size=15, color=WHITE)
    add_text(slide, 7.15, 3.30, 5.5, 0.28, "DECISION LATENCY: SIGNAL TO ACTION",
             "Roboto Medium", 11, LILAC)
    steps = ["Signal", "Investigation", "Decision", "Action"]
    for i, step in enumerate(steps):
        x = 7.15 + (i % 4) * 1.42
        add_rect(slide, x, 3.70, 1.30, 0.62, WHITE)
        add_text(slide, x, 3.70, 1.30, 0.62, step, "Roboto Medium", 11, DEEP,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 7.15, 4.55, 2.4, 0.2, "CURRENT", "Roboto Medium", 11, LILAC)
    add_text(slide, 7.15, 4.78, 2.6, 0.3, "Days to weeks", "Roboto Medium", 16, WHITE)
    add_text(slide, 10.1, 4.55, 2.4, 0.2, "TARGET", "Roboto Medium", 11, LILAC)
    add_text(slide, 10.1, 4.78, 2.6, 0.3, "Minutes", "Roboto Medium", 16, WHITE)
    return slide


def slide_modules(prs):
    slide = dark(prs, "Eight business decisions, supported by the same eight capabilities", 24)
    columns = [
        ("BUSINESS DECISIONS", [
            "Digital Funnel Performance",
            "Payment Performance Management",
            "Fraud Strategy Optimization",
            "Financial Crime Operations",
            "Dispute Operations Performance",
            "Liquidity and Funding Management",
            "Credit and Exposure Management",
            "Embedded Insights Strategy",
        ]),
        ("REUSABLE CAPABILITIES", [
            "Conversational Drill-Down",
            "Live Event Impact",
            "Transparent Answer",
            "Governed Business Context",
            "Scoped Access",
            "Embedded Streaming",
            "Structured Integration",
            "Reusable Decision View",
        ]),
    ]
    for col, (heading, items) in enumerate(columns):
        x = 0.64 + col * 6.35
        add_text(slide, x, 1.55, 5.8, 0.28, heading, "Roboto Medium", 12, LILAC)
        for i, item in enumerate(items):
            y = 2.00 + i * 0.52
            add_rect(slide, x, y, 0.42, 0.40, PURPLE)
            add_text(slide, x, y, 0.42, 0.40, str(i + 1), "Roboto Medium", 14, WHITE,
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            add_text(slide, x + 0.56, y, 5.3, 0.40, item, "Roboto", 15, WHITE,
                     anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 0.64, 6.30, 12.0, 0.32,
             "Each path connects:  Business decision  →  Capability  →  Role value  →  Validated delta.  Platform decisions: Analytics Governance, Analytics Integration Architecture, AI Platform Economics.",
             "Roboto", 13, LILAC)
    return slide


ROLES = [
    {
        "title": "Find where account applicants drop off, and why",
        "takeaway": "See Intent Earlier",
        "job": "Head of Digital Banking",
        "decision": "Digital Funnel Performance",
        "issue": "Improve digital account-opening conversion and act on customer intent before opportunities are lost.",
        "problems": [
            "Monthly reporting delays intervention",
            "Onboarding, behavioral, and relationship data stay disconnected",
            "Manual-review friction and affected cohorts are hard to isolate",
        ],
        "date": "By [campaign, product release, or planning date], because findings must inform the next journey or investment decision.",
        "capabilities": [
            "Ask the funnel by step and channel in natural language, instead of waiting on a monthly report",
            "Compare straight-through and manual review in one answer, so the friction shows up immediately",
            "Isolate the abandonment cohort in the conversation, rather than requesting another extract",
            "Return governed tables and charts the team can act on in the same session",
        ],
        "deltas": [
            ("Question-to-answer time", "[baseline]  →  [target]"),
            ("Report requests", "[baseline]  →  [target]"),
            ("Completion rate", "[baseline]%  →  [target]%"),
        ],
    },
    {
        "title": "Spot payment problems that cost revenue before customers notice",
        "takeaway": "Protect Every Payment",
        "job": "Head of Payments",
        "decision": "Payment Performance Management",
        "issue": "Protect payment revenue and customer experience by keeping approval rates and latency within thresholds.",
        "problems": [
            "Batch reports expose changes too late",
            "Processor, rail, merchant, region, and decline data are fragmented",
            "Operations cannot quickly isolate the affected window and segment",
        ],
        "date": "By [peak-volume event, processor migration, or contract renewal], because performance must be validated before the commercial commitment.",
        "capabilities": [
            "Show hourly approval, latency, and declines as they move, not in the next batch",
            "Compare the affected window with baseline in the same answer",
            "Rank the processor or rail behind the change without a new data pull",
            "Drill to merchant, region, and response code before customers notice",
            "Show the query behind the result; Aura does not move or retry the payment",
        ],
        "deltas": [
            ("Approval rate", "[baseline]%  →  [target]%"),
            ("Revenue at risk", "$[baseline]  →  $[target]"),
            ("Investigation time", "[baseline] hrs  →  [target] min"),
        ],
    },
    {
        "title": "Reduce losses without blocking good customers",
        "takeaway": "Trace Loss Faster",
        "job": "VP of Fraud Strategy",
        "decision": "Fraud Strategy Optimization",
        "issue": "Reduce fraud loss while controlling false positives and preserving legitimate revenue.",
        "problems": [
            "Fraud-type reporting obscures cross-channel patterns",
            "Investigators lack merchant, device, account-age, and velocity context",
            "Manual segmentation delays response",
        ],
        "date": "By [seasonal peak, product launch, or fraud-strategy review], because controls must be validated before exposure increases.",
        "capabilities": [
            "Compare loss and false positives by segment in one pass, not a report cycle",
            "Drill channel, merchant, device, age, and velocity without hand-built segments",
            "Hand investigators a ready cohort instead of a multi-day segmentation project",
            "Show the query and source rows so the pattern can be checked immediately",
            "Leave the block-or-allow decision with the investigator",
        ],
        "deltas": [
            ("Fraud loss", "$[baseline]  →  $[target]"),
            ("False positives", "[baseline]%  →  [target]%"),
            ("Pattern-isolation time", "[baseline]  →  [target]"),
        ],
    },
    {
        "title": "Spend reviewer time on decisions, not on gathering evidence",
        "takeaway": "Assemble Evidence Faster",
        "job": "Head of AML Operations",
        "decision": "Financial Crime Operations",
        "issue": "Clear alerts within policy while improving review quality and preserving human disposition.",
        "problems": [
            "Reviewers manually gather transactions, entities, counterparties, history",
            "Aging backlogs obscure consequential cohorts",
            "Fragmented evidence raises effort and inconsistent escalation",
        ],
        "date": "By [regulatory examination, policy review, or backlog milestone], because operations must show policy adherence and sufficient capacity.",
        "capabilities": [
            "Read backlog by age, rule, risk band, and queue without building a spreadsheet",
            "Prioritize the cohort with the team's own criteria in the same session",
            "Assemble transactions, entities, counterparties, and history in one answer",
            "Keep the packet scoped to the reviewer and inspectable for the exam file",
            "Deliver the evidence only; disposition stays with the reviewer",
        ],
        "deltas": [
            ("Review time", "[baseline]  →  [target]"),
            ("Median alert age", "[baseline]  →  [target]"),
            ("Cases per reviewer", "[baseline]  →  [target]"),
        ],
    },
    {
        "title": "Catch disputes before they miss their deadline",
        "takeaway": "Resolve Cases Earlier",
        "job": "VP of Disputes Operations",
        "decision": "Dispute Operations Performance",
        "issue": "Resolve disputes within SLA at lower cost while protecting customer trust.",
        "problems": [
            "Queue, transaction, contact, and credit data sit in separate systems",
            "Totals do not explain breach risk",
            "Agents manually reconstruct customer and case timelines",
        ],
        "date": "By [peak dispute period, SLA review, or regulatory milestone], because capacity must be in place before case volume increases.",
        "capabilities": [
            "Surface the disputes closest to an SLA breach before the clock runs out",
            "Compare reason, age, volume, and workload without exporting the queue",
            "Join transaction, contact, credit, and case events into one timeline",
            "Reuse that timeline on the next touch instead of rebuilding it",
            "Leave resolution and credit decisions with the authorized agent",
        ],
        "deltas": [
            ("Handle time", "[baseline]  →  [target]"),
            ("SLA breaches", "[baseline]%  →  [target]%"),
            ("Cost per case", "$[baseline]  →  $[target]"),
        ],
    },
    {
        "title": "See cash shortfalls while there is still time to fund them",
        "takeaway": "See Funding Pressure",
        "job": "Corporate Treasurer",
        "decision": "Liquidity and Funding Management",
        "issue": "Maintain sufficient liquidity and meet settlement obligations at the lowest practical funding cost.",
        "problems": [
            "Positions, settlements, queued payments, forecasts use separate reports",
            "Current and expected cash diverge without explanation",
            "Reconciliation delays conceal funding pressure",
        ],
        "date": "By [peak settlement cycle, funding window, or reporting date], because treasury needs time to act before obligations become exceptions.",
        "capabilities": [
            "Show fresh positions by entity and currency, not a stale pack",
            "Rank forecast variances in the same answer",
            "List payments due inside the next funding window",
            "Test a new event's liquidity impact without a new model",
            "Show the pressure only; treasury still releases or stops payments",
        ],
        "deltas": [
            ("Position freshness", "[baseline]  →  [target]"),
            ("Reconciliation time", "[baseline]  →  [target]"),
            ("Emergency funding cost", "$[baseline]  →  $[target]"),
        ],
    },
    {
        "title": "Give credit reviewers the full picture before they decide",
        "takeaway": "Context Before Decision",
        "job": "Chief Credit Officer",
        "decision": "Credit and Exposure Management",
        "issue": "Make timely, policy-compliant credit decisions while controlling concentration risk.",
        "problems": [
            "Applicant evidence and portfolio exposure reside in separate tools",
            "Missing evidence and manual review increase backlog",
            "Concentration by industry, geography, and risk grade is hard to assess",
        ],
        "date": "By [portfolio review, policy update, examination, or planning date], because leadership must show timely decisions and controlled exposure.",
        "capabilities": [
            "Show volume, backlog, review rate, and decision time without a status pack",
            "Flag missing evidence before the file sits in the queue",
            "Compare requested exposure with current concentration in one view",
            "Apply Domain context and scoped access so the packet is ready to review",
            "Leave approval, decline, pricing, and policy with the credit officer",
        ],
        "deltas": [
            ("Review time", "[baseline]  →  [target]"),
            ("Missing-evidence cases", "[baseline]%  →  [target]%"),
            ("Concentration off policy", "[baseline]  →  [target]"),
        ],
    },
    {
        "title": "Explain customers' own money to them, inside your app",
        "takeaway": "Explain Engage Monetize",
        "job": "Chief Product Officer",
        "decision": "Embedded Insights Strategy",
        "issue": "Increase digital engagement and monetization through trusted, personalized financial explanations.",
        "problems": [
            "Customers lack context for changes in financial activity",
            "Support repeatedly answers common analytical questions",
            "Embedded experiences need scoped data, continuity, security, disclosures",
        ],
        "date": "By [digital release or premium-product launch], because security, disclosures, experience design, and measurement must be validated before rollout.",
        "capabilities": [
            "Embed Aura Analyst through /chat with SSE streaming",
            "Continue contextual multi-turn sessions",
            "Restrict answers to the authenticated customer's Domain",
            "Support explanations with relevant data and charts",
            "Use /query for structured output; your team owns UI and actions",
        ],
        "deltas": [
            ("Engagement", "[baseline]  →  [target]"),
            ("Support deflection", "[baseline]%  →  [target]%"),
            ("Premium adoption", "[baseline]%  →  [target]%"),
        ],
    },
]


def slide_role(prs, role):
    slide = light(prs, role["title"], 22)
    add_text(slide, 9.85, 0.42, 2.85, 0.2, "TAKEAWAY", "Roboto Medium", 11, GRAY)
    add_text(slide, 9.85, 0.64, 2.85, 0.7, role["takeaway"], "Roboto Medium", 16, PURPLE)
    hline(slide, 0.64, 1.48, 12.05)
    label(slide, 0.64, 1.64, 6.1, "JOB TITLE")
    add_text(slide, 0.64, 1.86, 6.1, 0.32, role["job"], "Roboto Medium", 16, INK)
    label(slide, 0.64, 2.28, 6.1, role["decision"].upper())
    add_text(slide, 0.64, 2.50, 6.1, 0.72, role["issue"], "Roboto", 13, INK)
    label(slide, 0.64, 3.26, 6.1, "PROBLEMS / REASONS")
    bullets(slide, 0.64, 3.48, 6.1, 1.35, role["problems"], size=13)
    label(slide, 0.64, 4.95, 6.1, "CRITICAL DATE")
    add_text(slide, 0.64, 5.17, 6.1, 1.35, role["date"], "Roboto", 13, INK)
    label(slide, 7.15, 1.64, 5.5, "Specific Capabilities")
    bullets(slide, 7.15, 1.88, 5.5, 2.35, role["capabilities"], size=13)
    label(slide, 7.15, 4.35, 5.5, "DELTA")
    for i, (name, value) in enumerate(role["deltas"]):
        y = 4.62 + i * 0.68
        add_rect(slide, 7.15, y, 5.5, 0.60, MIST)
        add_text(slide, 7.30, y, 5.2, 0.28, name, "Roboto Medium", 12, INK, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, 7.30, y + 0.26, 5.2, 0.30, value, "Roboto Medium", 13, PURPLE)
    return slide


def slide_roles_overview(prs):
    slide = light(prs, "Each role asks a different question and needs different evidence", 24)
    lanes = [
        (PURPLE, "Business leaders",
         "Prioritize growth, conversion, margin, liquidity, and decision speed.",
         "Key capabilities: drill-down, live impact, decision views",
         "Success signal: faster, higher-value decisions"),
        (BLUE, "Risk leaders",
         "Prioritize evidence, policy context, access scope, explainability, and accountable ownership.",
         "Key capabilities: transparency, context, scoped access",
         "Success signal: defensible, inspectable decisions"),
        (ORANGE, "Technology leaders",
         "Prioritize data paths, latency, concurrency, integration contracts, and operational boundaries.",
         "Key capabilities: streaming, structured integration",
         "Success signal: feasible, ownable integration"),
    ]
    for i, (color, name, body, caps, success) in enumerate(lanes):
        x = 0.64 + i * 4.15
        add_rect(slide, x, 1.60, 3.95, 4.55, MIST)
        add_rect(slide, x, 1.60, 3.95, 0.58, color)
        add_text(slide, x + 0.18, 1.60, 3.6, 0.58, name, "Roboto Medium", 16, WHITE,
                 anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, x + 0.18, 2.40, 3.6, 1.35, body, "Roboto", 14, INK)
        add_text(slide, x + 0.18, 3.90, 3.6, 0.9, caps, "Roboto Medium", 13, color)
        add_text(slide, x + 0.18, 4.90, 3.6, 0.9, success, "Roboto", 13, GRAY)
    add_text(slide, 0.64, 6.30, 12.05, 0.5,
             "We focus on one primary role and up to two supporting roles. Domain and access views appear when they show how context and permissions shape the decision.",
             "Roboto", 13, INK)
    return slide


def slide_capabilities(prs):
    slide = light(prs, "Eight capabilities you can see working and test for yourself", 24)
    groups = [
        ("INTERACTION", [
            ("1", "Conversational Drill-Down", "Investigative speed"),
            ("2", "Live Event Impact", "Responsiveness"),
        ]),
        ("TRUST", [
            ("3", "Transparent Answer", "Inspectability, not guaranteed correctness"),
            ("4", "Governed Business Context", "Configured business context"),
        ]),
        ("INTEGRATION", [
            ("5", "Scoped Access", "Configured permissions"),
            ("6", "Embedded Streaming", "/chat, SSE, sessions"),
        ]),
        ("REUSE", [
            ("7", "Structured Integration", "/query, structured JSON"),
            ("8", "Reusable Decision View", "Repeat use and validation"),
        ]),
    ]
    for i, (heading, items) in enumerate(groups):
        x = 0.64 + i * 3.15
        add_text(slide, x, 1.60, 2.95, 0.28, heading, "Roboto Medium", 12, GRAY)
        hline(slide, x, 1.95, 2.95, PURPLE)
        for j, (num, name, detail) in enumerate(items):
            y = 2.20 + j * 2.15
            add_rect(slide, x, y, 2.95, 1.95, MIST)
            add_text(slide, x + 0.16, y + 0.16, 2.6, 0.36, num, "Roboto Medium", 20, PURPLE)
            add_text(slide, x + 0.16, y + 0.58, 2.6, 0.7, name, "Roboto Medium", 14, INK)
            add_text(slide, x + 0.16, y + 1.30, 2.6, 0.5, detail, "Roboto", 12, GRAY)
    return slide


def slide_tell(prs):
    slide = light(prs, "Every demo states the problem, shows the answer, and names what changed", 24)
    steps = [
        ("Tell", "State the decision, baseline time-to-know, consequence, and hypothesized delta."),
        ("Show", "Follow one decision path using only the capabilities needed to reach evidence."),
        ("Tell", "State what changed, who acts, and which of your metrics validates value."),
    ]
    for i, (name, body) in enumerate(steps):
        x = 0.64 + i * 4.15
        add_rect(slide, x, 1.60, 3.95, 2.15, MIST)
        add_text(slide, x + 0.2, 1.74, 3.55, 0.4, name, "Roboto Medium", 20, PURPLE)
        add_text(slide, x + 0.2, 2.20, 3.55, 1.35, body, "Roboto", 14, INK)
    label(slide, 0.64, 4.05, 6.0, "MEASURE THE BEFORE-AND-AFTER")
    bullets(slide, 0.64, 4.32, 6.0, 1.8, [
        "Investigation time and handoffs",
        "Evidence retrieval effort",
        "Decision-cycle duration",
        "Avoided rework",
    ], size=14)
    label(slide, 7.15, 4.05, 5.5, "KEEP THE STORY TIGHT")
    add_text(slide, 7.15, 4.32, 5.5, 1.8,
             "We remove anything that does not strengthen the decision-to-delta story. Every capability shown must move the before-and-after time gauge.",
             "Roboto", 14, INK)
    return slide


def slide_trust(prs):
    slide = light(prs, "You can trust an answer when you can check it and a person owns the decision", 22)
    items = [
        ("Domain context", "Make weak or missing Domain context visible rather than implying universal accuracy."),
        ("Scoped access", "Demonstrate scoped access rather than claiming universal governance; it depends on configuration, identity, policy, and deployment."),
        ("Inspectable evidence", "Use inspectability to support review without presenting it as guaranteed correctness."),
        ("Human ownership", "Keep regulated decisions with accountable humans; Preview capabilities are not autonomous controls."),
    ]
    for i, (name, body) in enumerate(items):
        y = 1.60 + i * 1.22
        add_rect(slide, 0.64, y, 0.10, 1.05, PURPLE)
        add_text(slide, 1.00, y, 11.4, 0.36, name, "Roboto Medium", 18, INK)
        add_text(slide, 1.00, y + 0.40, 11.4, 0.6, body, "Roboto", 14, GRAY)
    return slide


def slide_apis(prs):
    slide = light(prs, "Aura supplies the answers; your team builds the experience around them", 22)
    panels = [
        ("Native Aura intelligence", [
            "/chat provides conversation, SSE streaming, and session continuity",
            "/query provides structured JSON for downstream rendering and integration",
            "Scope includes Aura Analyst and API response patterns",
            "These capabilities remain Preview",
        ]),
        ("Your team's experience", [
            "Custom UI and workflow actions",
            "Telemetry and downstream orchestration",
            "Security, disclosures, and approved actions",
            "Engineering, product, security, and operational work beyond the intelligence layer",
        ]),
    ]
    for i, (heading, items) in enumerate(panels):
        x = 0.64 + i * 6.35
        add_rect(slide, x, 1.60, 6.05, 3.55, MIST)
        add_rect(slide, x, 1.60, 6.05, 0.08, PURPLE if i == 0 else BLUE)
        add_text(slide, x + 0.24, 1.84, 5.6, 0.4, heading, "Roboto Medium", 16, INK)
        bullets(slide, x + 0.24, 2.40, 5.6, 2.5, items, size=14)
    add_rect(slide, 0.64, 5.35, 12.05, 1.20, DEEP)
    add_text(slide, 0.88, 5.48, 11.6, 0.28, "API BOUNDARY", "Roboto Medium", 11, LILAC)
    add_text(slide, 0.88, 5.80, 11.6, 0.6,
             "Embedded delivery is a joint effort: native intelligence plus your engineering, product, security, and operations.",
             "Roboto", 15, WHITE)
    return slide


def slide_formats(prs):
    slide = light(prs, "Sessions from 5 to 40 minutes tell the same story at different depths", 24)
    formats = [
        ("5 min", "One role, one decision, one capability, and one delta hypothesis."),
        ("10 min", "One decision, two capabilities, Tell–Show–Tell, and one measurable delta."),
        ("15 min", "Primary and supporting roles, one decision, three capabilities, and the native-versus-custom boundary."),
        ("25 min", "Two related decisions, all three role perspectives, an API path, guardrails, and a validation plan."),
        ("40 min", "Discovery-led design, architecture review, ownership mapping, delta validation, and a workshop checkpoint."),
    ]
    for i, (time, body) in enumerate(formats):
        x = 0.64 + i * 2.50
        add_rect(slide, x, 1.65, 2.35, 3.35, MIST)
        add_rect(slide, x, 1.65, 2.35, 0.08, PURPLE)
        add_text(slide, x + 0.14, 1.90, 2.08, 0.55, time, "Roboto Medium", 20, PURPLE)
        add_text(slide, x + 0.14, 2.55, 2.08, 2.2, body, "Roboto", 13, INK)
    add_text(slide, 0.64, 5.25, 12.05, 0.55,
             "Depth and audience coverage increase from left to right; the decision story stays constant.",
             "Roboto", 14, GRAY)
    return slide


def slide_combinations(prs):
    slide = dark(prs, "Suggested pairings for common fintech priorities", 24)
    groups = [
        ("Growth", ["Digital Funnel Performance", "Conversational Drill-Down", "Reusable Decision View"]),
        ("Payments", ["Payment Performance Management", "Live Event Impact", "Embedded Streaming"]),
        ("Financial crime", ["Fraud Strategy Optimization / Financial Crime Operations", "Transparent Answer", "Scoped Access"]),
        ("Operations", ["Dispute Operations Performance", "Governed Business Context", "Reusable Decision View"]),
        ("Embedded product", ["Embedded Insights Strategy", "Embedded Streaming", "Structured Integration"]),
    ]
    for i, (heading, items) in enumerate(groups):
        x = 0.64 + i * 2.50
        add_text(slide, x, 1.70, 2.35, 0.7, heading, "Roboto Medium", 16, LILAC)
        for j, item in enumerate(items):
            y = 2.60 + j * 1.15
            add_rect(slide, x, y, 2.35, 1.00, DEEP)
            add_text(slide, x + 0.12, y, 2.12, 1.00, item, "Roboto", 13, WHITE,
                     anchor=MSO_ANCHOR.MIDDLE)
    return slide


def slide_questions(prs):
    slide = light(prs, "Six questions to decide whether this is worth pursuing together", 24)
    questions = [
        ("1", "Job Title", "Confirm one specific job title, not a department or broad role."),
        ("2", "Critical Business Issue", "Validate the critical business issue and the problems preventing progress."),
        ("3", "Problems", "Name the reasons progress stalls today."),
        ("4", "Capabilities", "Focus only on the capabilities you consider must-have."),
        ("5", "Delta", "Quantify with your baselines, targets, and methods."),
        ("6", "Critical Date", "Tie to a consequential business, regulatory, product, or operating event."),
    ]
    for i, (num, name, body) in enumerate(questions):
        x = 0.64 + (i % 6) * 2.08
        add_text(slide, x, 1.70, 1.95, 0.45, num, "Roboto Medium", 28, PURPLE)
        add_text(slide, x, 2.30, 1.95, 0.85, name, "Roboto Medium", 14, INK)
        add_text(slide, x, 3.25, 1.95, 2.4, body, "Roboto", 13, GRAY)
    add_text(slide, 0.64, 6.05, 12.05, 0.5,
             "Each link must hold before we commit to next steps.", "Roboto Medium", 16, INK)
    return slide


def slide_gate(prs):
    slide = light(prs, "Go ahead, change course, or stop, based on value you can test", 24)
    inputs = [
        ("Decision", "Priority decision, primary job title, decision owner, consequential question"),
        ("Data", "Data and events, Domain context, API fit"),
        ("Controls", "Access, engineering needs, ownership boundaries, governance dependencies"),
        ("Delta", "Minimum capabilities, your success evidence, delta assumptions"),
    ]
    for i, (name, body) in enumerate(inputs):
        x = 0.64 + i * 3.15
        add_rect(slide, x, 1.60, 3.00, 2.15, MIST)
        add_text(slide, x + 0.16, 1.74, 2.7, 0.36, name, "Roboto Medium", 16, PURPLE)
        add_text(slide, x + 0.16, 2.18, 2.7, 1.35, body, "Roboto", 13, INK)
    add_text(slide, 0.64, 4.00, 4, 0.3, "DECISION GATE", "Roboto Medium", 12, GRAY)
    outcomes = [
        (PURPLE, "Proceed", "to validation"),
        (BLUE, "Reshape", "the hypothesis"),
        (GRAY, "Stop", "and revisit later"),
    ]
    for i, (color, name, detail) in enumerate(outcomes):
        x = 0.64 + i * 4.15
        add_rect(slide, x, 4.40, 3.95, 1.15, color)
        add_text(slide, x + 0.2, 4.50, 3.55, 0.45, name, "Roboto Medium", 18, WHITE)
        add_text(slide, x + 0.2, 4.98, 3.55, 0.4, detail, "Roboto", 14, WHITE)
    add_text(slide, 0.64, 5.80, 12.05, 0.7,
             "Record inspectability limits, human accountability, and Preview status before deciding.",
             "Roboto", 14, GRAY)
    return slide


def delete_all_slides(prs):
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        rel_id = slide_id.get(qn("r:id"))
        prs.part.drop_rel(rel_id)
        slide_ids.remove(slide_id)


def main():
    source = Presentation(SOURCE)
    shutil.copy(TEMPLATE, DEST)
    prs = Presentation(DEST)
    delete_all_slides(prs)
    built = [
        slide_cover(prs),
        slide_modules(prs),
        *[slide_role(prs, role) for role in ROLES],
        slide_roles_overview(prs),
        slide_capabilities(prs),
        slide_tell(prs),
        slide_trust(prs),
        slide_apis(prs),
        slide_formats(prs),
        slide_combinations(prs),
        slide_questions(prs),
        slide_gate(prs),
    ]
    assert len(built) == len(source.slides) == 19
    for slide, original in zip(built, source.slides):
        copy_note(slide, original)
    prs.core_properties.title = "Aura Analyst for Fintech: Decision Playbook"
    prs.save(DEST)
    print(f"Wrote {DEST} ({len(built)} slides)")


if __name__ == "__main__":
    main()
